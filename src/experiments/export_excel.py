"""
export_excel.py — rangkum seluruh hasil run jadi SATU workbook Excel untuk dokumentasi.

Membaca (yang tidak ada dilewati; statusnya ditulis di sheet "Info"):
    data/raw_logs/{baseline,graphrag}_log.json      -> Overall, By Difficulty, Retrieval, SQL Diagnostics, Per Query
    outputs/tables/confound_diagnostics.json        -> SQL Diagnostics (pipeline.py --baseline / eval_pipeline.ipynb)
    outputs/tables/spider_official_eval.csv         -> Spider Official (eval_pipeline.ipynb bagian 14)
    outputs/tables/dimensions_results.json          -> Dimensions, D1..D6 (run_all_dimensions.py)
    outputs/tables/ablation_results.csv             -> Ablation (few-shot k, tahap 2)
    outputs/tables/model_comparison.csv             -> Models (tahap 3)
    outputs/tables/sla_trial_recap.csv, sla_trial_T*_K*.csv -> SLA Trial (graphrag_sla_trial.ipynb)

Tabelnya dibangun oleh src/experiments/result_tables.py, modul yang SAMA dengan yang ditampilkan
eval_pipeline.ipynb, jadi angka notebook dan workbook identik. Isinya hanya angka faktual:
teks interpretasi/kesimpulan dari dimensions_results.json tidak ikut diekspor.

Setiap tabel adalah Excel Table (filter + baris berselang), persentase disimpan sebagai pecahan
dengan format %, selisih persentase dalam poin persen (pp).

Usage (dari root repo, lokal maupun Kaggle):
    python src/experiments/export_excel.py
    python src/experiments/export_excel.py --out outputs/tables/hasil_skripsi.xlsx
    python src/experiments/export_excel.py --skip-missing --out hasil_trial.xlsx   # ekspor parsial
    python src/experiments/export_excel.py --raw-logs-dir data/raw_logs/Qwen2.5-Coder-3B-Instruct \
        --tables-dir outputs/models/Qwen2.5-Coder-3B-Instruct/tables
"""

import argparse
import glob
import json
import re
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from src.experiments import result_tables as rt  # noqa: E402

# ---------------------------------------------------------------------------
# Style
# ---------------------------------------------------------------------------

FONT = "Arial"
NAVY, GREY = "1F3864", "595959"
TABLE_STYLE = "TableStyleMedium2"
TEXT_WIDTH_CAP = 60      # kolom teks panjang (SQL) dipotong lebarnya, tidak di-wrap
NUM_WIDTH = 12


def _font(bold=False, italic=False, color="000000", size=10):
    return Font(name=FONT, bold=bold, italic=italic, color=color, size=size)


def _cell_value(v):
    if rt.is_missing(v):
        return None
    if hasattr(v, "item"):   # numpy scalar -> Python
        v = v.item()
    return v


class Sheet:
    """Satu worksheet: judul, catatan, lalu satu atau lebih tabel (masing-masing Excel Table)."""

    _n_tables = 0   # nama Excel Table harus unik di seluruh workbook

    def __init__(self, wb: Workbook, name: str, title: str, note: str | None = None, first: bool = False):
        self.ws = wb.active if first else wb.create_sheet()
        self.ws.title = name
        self.ws.sheet_view.showGridLines = False
        self.ws.page_setup.orientation = "landscape"
        self.ws.page_setup.fitToWidth = 1
        self.ws.page_setup.fitToHeight = 0
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.widths: dict[int, float] = {}
        self.header_rows: list[int] = []
        self.has_data = False
        self.r = 1
        self.ws.cell(self.r, 1, title).font = _font(bold=True, size=14, color=NAVY)
        self.r += 1
        if note:
            self.note(note)
        self.r += 1

    def note(self, text: str):
        self.ws.cell(self.r, 1, text).font = _font(italic=True, color=GREY)
        self.r += 1

    def _fit(self, col: int, width: float):
        self.widths[col] = max(self.widths.get(col, 8), min(width, TEXT_WIDTH_CAP))

    def table(self, df: pd.DataFrame, caption: str | None = None):
        caption = caption or df.attrs.get("title")
        if caption:
            self.ws.cell(self.r, 1, caption).font = _font(bold=True, size=11, color=NAVY)
            self.r += 1
        cols = rt.visible_columns(df)
        if df.empty or not cols:
            self.note("No rows.")
            self.r += 1
            return
        header = self.r
        for j, col in enumerate(cols, 1):
            c = self.ws.cell(header, j, str(col))
            c.font = _font(bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            self._fit(j, len(str(col)) * 1.1 + 4)
        self.ws.row_dimensions[header].height = 30
        for i in range(len(df)):
            row = header + 1 + i
            for j, col in enumerate(cols, 1):
                v = _cell_value(df[col].iloc[i])
                c = self.ws.cell(row, j, v)
                c.font = _font()
                kind = rt.cell_kind(df, i, col)
                xl = rt.KINDS.get(kind, rt.KINDS["text"])[1]
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    if xl:
                        c.number_format = xl
                    c.alignment = Alignment(horizontal="right")
                    self._fit(j, NUM_WIDTH)
                else:
                    if isinstance(v, bool):
                        c.value = "Yes" if v else "No"
                    c.alignment = Alignment(horizontal="left")
                    self._fit(j, len(str(v or "")) * 1.05 + 2)
        last = header + len(df)
        Sheet._n_tables += 1
        ref = f"A{header}:{get_column_letter(len(cols))}{last}"
        tbl = Table(displayName=f"T{Sheet._n_tables:03d}_{re.sub(r'[^A-Za-z0-9]', '', self.ws.title)[:20]}", ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(name=TABLE_STYLE, showRowStripes=True)
        self.ws.add_table(tbl)
        self.header_rows.append(header)
        self.has_data = True
        self.r = last + 2

    def finish(self, freeze_col: str = "A"):
        for col, w in self.widths.items():
            self.ws.column_dimensions[get_column_letter(col)].width = w
        if len(self.header_rows) == 1:   # satu tabel: header tetap terlihat saat scroll
            self.ws.freeze_panes = f"{freeze_col}{self.header_rows[0] + 1}"


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------

def _read_json(p: Path):
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def _read_csv(p: Path, **kw):
    try:
        df = pd.read_csv(p, **kw)
        return df if len(df) else None
    except (OSError, ValueError, pd.errors.EmptyDataError):
        return None


def _alpha_default() -> float:
    try:
        from src.core.config import PipelineConfig
        return float(PipelineConfig().token_output_weight)
    except Exception:   # torch tidak ada -> pakai nilai tetap dari proposal
        return 3.0


DEFINITIONS = [
    ("EX", "Execution Accuracy: share of queries whose predicted SQL returns the same result as the gold SQL "
           "(official Spider eval_exec_match, order-sensitive)."),
    ("ESM", "Exact Set Match: share of queries whose parsed SQL structure matches the gold SQL."),
    ("CM", "Component Match: per-clause match rate; UNION ALL is not distinguishable by the Spider parser and stays empty."),
    ("SLA", "Schema Linking Accuracy: precision/recall/F1 of the schema sent to the LLM against the schema used by the "
            "gold SQL, macro-averaged per query (src/metrics/sla.py). Queries with an empty gold schema score 0."),
    ("T", "Token consumption per query: T = T_in + α × T_out."),
    ("TEP", "Token Elasticity of Performance: (ΔEX / EX_Baseline) / (ΔT / T_Baseline)."),
    ("QVT", "Query Variance Testing over natural paraphrase pairs in dev.json; groups with every variant wrong are excluded."),
    ("Δ", "GraphRAG − Baseline. For percentages the difference is in percentage points (pp)."),
    ("Relative change", "(GraphRAG − Baseline) / Baseline."),
]


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def build_workbook(raw_logs_dir: Path, tables_dir: Path, out: Path, skip_missing: bool = False) -> Path:
    """skip_missing=True: buang sheet yang sumbernya tidak ada (ekspor parsial, mis. hanya trial SLA
    atau hanya ablation), alih-alih menampilkan sheet berisi catatan "source not found"."""
    Sheet._n_tables = 0
    alpha = _alpha_default()

    logs = {c: _read_json(raw_logs_dir / f"{c}_log.json") for c in ("baseline", "graphrag")}
    dims = _read_json(tables_dir / "dimensions_results.json")
    confound = _read_json(tables_dir / "confound_diagnostics.json")
    spider = _read_csv(tables_dir / "spider_official_eval.csv")
    ablation = _read_csv(tables_dir / "ablation_results.csv")
    models = _read_csv(tables_dir / "model_comparison.csv")
    trial = _read_csv(tables_dir / "sla_trial_recap.csv")
    cfg_frames = [d for f in sorted(glob.glob(str(tables_dir / "sla_trial_T*_K*.csv")))
                  if (d := _read_csv(Path(f), index_col=0)) is not None]
    have_logs = bool(logs["baseline"]) and bool(logs["graphrag"])
    b, g = logs["baseline"] or [], logs["graphrag"] or []

    wb = Workbook()
    info = Sheet(wb, "Info", "Text-to-SQL evaluation results: GraphRAG vs Baseline (Spider dev)",
                 f"Generated by src/experiments/export_excel.py on {datetime.now():%Y-%m-%d %H:%M}. "
                 "All values are read from run output files.", first=True)
    sources = [
        ("raw_logs Baseline", logs["baseline"] is not None, raw_logs_dir / "baseline_log.json"),
        ("raw_logs GraphRAG", logs["graphrag"] is not None, raw_logs_dir / "graphrag_log.json"),
        ("Confound diagnostics", confound is not None, tables_dir / "confound_diagnostics.json"),
        ("Official Spider evaluation", spider is not None, tables_dir / "spider_official_eval.csv"),
        ("Dimensions 1-6", dims is not None, tables_dir / "dimensions_results.json"),
        ("Ablation stage 2 (few-shot k)", ablation is not None, tables_dir / "ablation_results.csv"),
        ("Model comparison (stage 3)", models is not None, tables_dir / "model_comparison.csv"),
        ("SLA trial (recap)", trial is not None, tables_dir / "sla_trial_recap.csv"),
        ("SLA trial (per threshold pair)", bool(cfg_frames), tables_dir / "sla_trial_T*_K*.csv"),
    ]
    info.table(pd.DataFrame([{"Source": s, "Found": "Yes" if ok else "No", "Path": str(p)} for s, ok, p in sources]),
               "Data sources")
    info.table(rt.with_formats(pd.DataFrame([{"Parameter": "α (output-token weight in T)", "Value": alpha},
                                              {"Parameter": "Queries in raw_logs (Baseline)", "Value": len(b)},
                                              {"Parameter": "Queries in raw_logs (GraphRAG)", "Value": len(g)}]),
                                {"Value": "num1"}), "Run parameters")
    info.table(pd.DataFrame(DEFINITIONS, columns=["Term", "Definition"]), "Definitions")
    info.ws.column_dimensions["B"].width = TEXT_WIDTH_CAP
    sheets = [info]

    def add(name: str, title: str, tables: list, missing: str, note: str | None = None, freeze_col: str = "A"):
        sh = Sheet(wb, name, title, note)
        for t in tables:
            sh.table(t)
        if not sh.has_data:
            sh.note(f"Source not found: {missing}")
        sh.finish(freeze_col)
        sheets.append(sh)

    raw_missing = f"{raw_logs_dir}/baseline_log.json and graphrag_log.json"
    add("Overall", "Overall comparison", [rt.overall_comparison(b, g, alpha)] if have_logs else [], raw_missing,
        f"Computed from raw_logs. T = T_in + {alpha:g} × T_out.")
    add("By Difficulty", "Metrics by difficulty", [rt.by_difficulty(b, g, alpha)] if have_logs else [], raw_missing)
    add("Retrieval", "Schema linking accuracy (SLA)",
        [rt.retrieval_metrics(b, g), rt.retrieval_by_difficulty(b, g)] if have_logs else [], raw_missing)
    diag = ([rt.prediction_checks(b, g), rt.outcome_matrix(b, g), rt.paired_ex(b, g), rt.sql_feature_usage(b, g),
             rt.cm_from_logs(b, g)] if have_logs else []) + ([rt.confound_table(confound)] if confound else [])
    add("SQL Diagnostics", "SQL diagnostics", diag, raw_missing,
        "SQL feature usage uses keyword matching on the SQL text. Confound counters are read-only and do not change any metric.")
    if spider is not None:
        parsed = {lab: sub.drop(columns="Pipeline") for lab, sub in spider.groupby("Pipeline", sort=False)}
        long = rt.with_formats(spider.copy(), {c: "num3" for c in spider.columns if c not in ("Pipeline", "Section", "Metric")},
                                title="Official Spider evaluation · all reported scores")
        add("Spider Official", "Official Spider evaluation (cross-check)", [rt.spider_summary(parsed), long], "")
    elif not skip_missing:
        add("Spider Official", "Official Spider evaluation (cross-check)", [],
            f"{tables_dir}/spider_official_eval.csv (written by eval_pipeline.ipynb section 14)")
    add("Per Query", "Per-query results", [rt.per_query(b, g, alpha)] if have_logs else [], raw_missing, freeze_col="D")

    add("Dimensions", "Analysis dimensions: status", [rt.dim_status(dims)] if dims else [], f"{tables_dir}/dimensions_results.json")
    for n, name in rt.DIM_NAMES.items():
        add(f"D{n}", f"Dimension {n}: {name}", rt.DIM_TABLES[n](dims), f"dimension_{n} with status 'ok' in dimensions_results.json")

    add("Ablation", "Ablation stage 2: few-shot k", rt.ablation_tables(ablation), f"{tables_dir}/ablation_results.csv")
    add("Models", "Ablation stage 3: LLM comparison", rt.model_tables(models), f"{tables_dir}/model_comparison.csv")
    add("SLA Trial", "GraphRAG schema-linking trial (SLA)", rt.sla_trial_tables(trial, cfg_frames),
        f"{tables_dir}/sla_trial_recap.csv")

    if skip_missing:
        for sh in sheets[1:]:
            if not sh.has_data:
                wb.remove(sh.ws)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description="Rangkum hasil run jadi satu file Excel.")
    ap.add_argument("--raw-logs-dir", default="data/raw_logs")
    ap.add_argument("--tables-dir", default="outputs/tables")
    ap.add_argument("--out", default="outputs/tables/hasil_skripsi.xlsx")
    ap.add_argument("--skip-missing", action="store_true",
                    help="Buang sheet yang sumber datanya belum ada (ekspor parsial: hanya trial SLA, hanya ablation, dst.).")
    args = ap.parse_args()
    out = build_workbook(Path(args.raw_logs_dir), Path(args.tables_dir), Path(args.out), args.skip_missing)
    print(f"Excel tersimpan: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
